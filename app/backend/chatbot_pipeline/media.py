"""
media.py - Speech, vision and document generation for the chat.

Four capabilities that all share one design rule: **every one of them is
optional, and the absence of a model is reported rather than crashed on.**

That rule is not politeness. These features need weights that are baked into
the container image at build time - a Whisper model, a Piper voice, a vision
LLM - and an air-gapped deployment cannot fetch a missing one at runtime. The
failure mode without capability checks is a 500 with a stack trace about a
missing file, on a machine with no internet, for a feature the operator may
not have known was optional. `capabilities()` below is what the UI reads to
decide whether to show the microphone at all.

    PPTX    python-pptx. Pure Python, no model, ~5MB. Always available if the
            package is installed.
    TTS     Piper via onnxruntime. Needs a .onnx voice + its .json config.
    STT     faster-whisper. Needs a CTranslate2 model directory.
    VISION  A multimodal model served by the same Ollama the chat already
            uses - no second runtime, no second port.

Model locations come from the environment so the Dockerfile can bake them
anywhere sensible, defaulting to /opt/models/<name>.
"""

import base64
import charts
import io
import logging
import re
import os
from typing import Any, Dict, List, Optional

logger = logging.getLogger(__name__)

PIPER_VOICE = os.getenv('PIPER_VOICE', '/opt/models/piper/en_US-lessac-medium.onnx')
WHISPER_MODEL = os.getenv('WHISPER_MODEL', '/opt/models/whisper-base')
WHISPER_COMPUTE = os.getenv('WHISPER_COMPUTE', 'int8')
# Beam search, not greedy. Measured on a domain sentence containing a hostname
# and an acronym: greedy scored 25% word error rate ("Host&DV and WNSXM GR01",
# "SC cluster"), beam 5 with the vocabulary hint below scored 10% - at the same
# transcription time. A larger model was the obvious fix and was the wrong one:
# `small` scored 15% and took 2.5x longer, so the win here is decoding, not
# parameters. Raise WHISPER_MODEL to small/medium only if accuracy still falls
# short after this.
WHISPER_BEAM = int(os.getenv('WHISPER_BEAM', '5'))

# Seeded vocabulary. Whisper conditions on this text, so naming the terms
# that actually occur in the deployment is what takes word error rate from
# 25% to 10% - an unusual hostname is otherwise transcribed as several
# English words, and ESXi as "SC".
#
# The DEFAULT below is generic on purpose. A deployment's real hostnames
# are exactly the terms worth seeding, and exactly the terms not to publish
# in source - so they go in .env (see .env.example), which ships to the
# operator rather than to a public repository.
WHISPER_PROMPT = os.getenv(
    'WHISPER_PROMPT',
    'Log analytics questions. Hostnames are short alphanumeric codes with '
    'hyphens, such as fw-edge-01 or srv-db-02. '
    'Severities ERROR WARNING CRITICAL NOTICE INFO DEBUG. '
    'Platforms ESXi, NSX, vCenter, vROps, Horizon, Kubernetes, CoreDNS, '
    'PostgreSQL, MinIO, Squid, auditd. Neo4j, Cypher, Kafka, SIEM, ULPF.')
# minicpm-v, not llava. The feature exists to read error text off a
# screenshot, and llava:7b cannot: asked to transcribe 30px black-on-white
# Consolas it answered "the image is too blurry", and on a smaller capture
# it invented an error about a host that was not in the picture. minicpm-v
# transcribed the same image exactly, hostname included. A model that
# fabricates hostnames inside a security tool is worse than no model.
VISION_MODEL = os.getenv('VISION_MODEL', 'minicpm-v')
# A 7B multimodal model on CPU needs minutes per image, not seconds -
# measured: llava:7b exceeded a 180s ceiling on a 760x260 screenshot. The
# chat timeout is sized for text and is far too short here.
VISION_TIMEOUT = int(os.getenv('VISION_TIMEOUT', '600'))
OLLAMA_HOST = os.getenv('OLLAMA_HOST', 'http://localhost:11434').rstrip('/')

# Bounded so a malformed or hostile request cannot exhaust memory. Speech is
# short by nature; an hour of audio is not a chat message.
MAX_AUDIO_BYTES = 25 * 1024 * 1024
MAX_IMAGE_BYTES = 12 * 1024 * 1024
MAX_TTS_CHARS = 4000


# ---------------------------------------------------------------------------
# Capability probing
# ---------------------------------------------------------------------------
def _has_module(name: str) -> bool:
    try:
        __import__(name)
        return True
    except Exception:
        return False


def _vision_available() -> bool:
    """True only if Ollama is reachable AND actually has the vision model.

    Checking the tag list rather than just the port matters: Ollama answers on
    its port whether or not the model was ever pulled, so a port check would
    advertise a capability that fails on first use - and on an air-gapped box
    it cannot be pulled to recover.
    """
    try:
        import requests
        r = requests.get(f"{OLLAMA_HOST}/api/tags", timeout=4)
        if r.status_code != 200:
            return False
        names = [m.get('name', '') for m in r.json().get('models', [])]
        base = VISION_MODEL.split(':')[0]
        return any(n == VISION_MODEL or n.split(':')[0] == base for n in names)
    except Exception:
        return False


def capabilities() -> Dict[str, Any]:
    """What this deployment can actually do, with the reason when it cannot."""
    piper_ok = _has_module('piper') and os.path.exists(PIPER_VOICE)
    whisper_ok = _has_module('faster_whisper') and os.path.isdir(WHISPER_MODEL)
    pptx_ok = _has_module('pptx')
    vision_ok = _vision_available()

    def why(ok, module, path, what):
        if ok:
            return None
        if not _has_module(module):
            return f"{module} is not installed in this image"
        return f"no {what} at {path}"

    return {
        'pptx': {'available': pptx_ok,
                 'detail': None if pptx_ok else 'python-pptx is not installed in this image'},
        'tts': {'available': piper_ok, 'model': PIPER_VOICE,
                'detail': why(piper_ok, 'piper', PIPER_VOICE, 'voice')},
        'stt': {'available': whisper_ok, 'model': WHISPER_MODEL,
                'detail': why(whisper_ok, 'faster_whisper', WHISPER_MODEL, 'model directory')},
        'vision': {'available': vision_ok, 'model': VISION_MODEL,
                   'detail': None if vision_ok
                   else f"Ollama at {OLLAMA_HOST} has no model matching {VISION_MODEL}"},
    }


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------
_MD_MARKS = re.compile(r'(\*\*|__|`)')


def _plain(text) -> str:
    """Strips inline markdown. A slide bullet and a Word bullet are plain
    text, so leaving the markers in prints literal asterisks into the
    document."""
    return _MD_MARKS.sub('', str(text or '')).strip()


def _add_bullets(frame, lines: List[str]) -> None:
    frame.clear()
    for i, line in enumerate(lines):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.text = _plain(line)
        para.level = 0


def build_pptx(question: str, answer: str, cypher: Optional[str] = None,
               rows: Optional[List[Dict[str, Any]]] = None) -> bytes:
    """One answer as a small deck: title, the answer, the query, the numbers.

    Deliberately plain. A generated deck that tries to art-direct itself ends
    up wrong on someone else's template; this produces something a human can
    paste into their own without undoing anything.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()

    # --- title -------------------------------------------------------------
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = 'Universal Log Pre-processing Framework'
    slide.placeholders[1].text = _plain(question or 'Chat answer')[:250]

    # --- the answer, split so a long one does not run off the slide --------
    paragraphs = [p.strip() for p in (answer or '').split('\n') if p.strip()]
    CHUNK = 12
    chunks = [paragraphs[i:i + CHUNK] for i in range(0, len(paragraphs), CHUNK)] or [['(no answer)']]
    for n, chunk in enumerate(chunks):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = 'Answer' + (f' ({n + 1}/{len(chunks)})' if len(chunks) > 1 else '')
        _add_bullets(s.placeholders[1].text_frame, chunk)

    # --- the query, so the number is auditable -----------------------------
    if cypher:
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = 'Query'
        box = s.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(8.8), Inches(4.5))
        tf = box.text_frame
        tf.word_wrap = True
        tf.text = cypher[:1800]
        for para in tf.paragraphs:
            for run in para.runs:
                run.font.size = Pt(11)
                run.font.name = 'Consolas'

    # --- the rows, as a chart and then as a table --------------------------
    # The chart comes first because it is what the reader looks at, and it is
    # inferred with the same rule the chat uses - so an exported answer carries
    # the chart the user was already looking at rather than only its numbers.
    chart = charts.from_rows(rows)
    if chart:
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = chart['title']
        charts.add_to_slide(s, chart, Inches(0.7), Inches(1.5), Inches(8.6), Inches(4.8))

    if rows:
        cols = list(rows[0].keys())[:6]
        shown = rows[:12]
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = f'Results ({len(rows)} row{"s" if len(rows) != 1 else ""})'
        table = s.shapes.add_table(len(shown) + 1, len(cols),
                                   Inches(0.5), Inches(1.6),
                                   Inches(9.0), Inches(0.4 * (len(shown) + 1))).table
        for c, name in enumerate(cols):
            table.cell(0, c).text = str(name)
        for r, row in enumerate(shown, start=1):
            for c, name in enumerate(cols):
                table.cell(r, c).text = str(row.get(name, ''))[:60]

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# TTS
# ---------------------------------------------------------------------------
_piper_voice = None


def synthesize(text: str) -> bytes:
    """Text to a WAV byte string, via Piper.

    The voice is loaded once and kept: it is ~60MB of ONNX and reloading it per
    request would make every spoken answer wait on disk for no reason.
    """
    global _piper_voice
    text = (text or '').strip()
    if not text:
        raise ValueError('nothing to speak')
    if len(text) > MAX_TTS_CHARS:
        text = text[:MAX_TTS_CHARS]

    from piper.voice import PiperVoice
    if _piper_voice is None:
        _piper_voice = PiperVoice.load(PIPER_VOICE)

    # synthesize_wav, not synthesize: the latter returns an iterable of audio
    # chunks and leaves the caller to write the RIFF header, which fails with
    # "# channels not specified" if you hand it a bare wave writer. The _wav
    # variant sets the format from the voice's own sample rate.
    import wave
    buf = io.BytesIO()
    with wave.open(buf, 'wb') as wav:
        _piper_voice.synthesize_wav(text, wav)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# STT
# ---------------------------------------------------------------------------
_whisper = None


def transcribe(audio: bytes, filename: str = 'audio.webm') -> Dict[str, Any]:
    """Audio bytes to text, via faster-whisper.

    int8 by default: the target VM is CPU-only and air-gapped, and int8 is
    roughly 4x faster than float32 for a barely perceptible accuracy cost on
    short dictation. Override with WHISPER_COMPUTE.
    """
    global _whisper
    if not audio:
        raise ValueError('no audio received')
    if len(audio) > MAX_AUDIO_BYTES:
        raise ValueError(f'audio too large ({len(audio)} bytes, limit {MAX_AUDIO_BYTES})')

    from faster_whisper import WhisperModel
    if _whisper is None:
        _whisper = WhisperModel(WHISPER_MODEL, device='cpu', compute_type=WHISPER_COMPUTE)

    import tempfile
    suffix = os.path.splitext(filename)[1] or '.webm'
    tmp = tempfile.NamedTemporaryFile(suffix=suffix, delete=False)
    try:
        tmp.write(audio)
        tmp.close()
        segments, info = _whisper.transcribe(
            tmp.name,
            beam_size=WHISPER_BEAM,
            initial_prompt=WHISPER_PROMPT,
            vad_filter=True,
            # A single dictated question has no previous text worth carrying,
            # and conditioning on it is what makes short clips loop a phrase.
            condition_on_previous_text=False,
        )
        text = ' '.join(seg.text.strip() for seg in segments).strip()
        return {'text': text, 'language': getattr(info, 'language', None),
                'duration': round(getattr(info, 'duration', 0.0), 2)}
    finally:
        try:
            os.unlink(tmp.name)
        except OSError:
            pass


# ---------------------------------------------------------------------------
# Vision
# ---------------------------------------------------------------------------
_VISION_SYSTEM = (
    "You are looking at a screenshot an operations engineer has pasted into a "
    "log analytics console - typically a terminal, a dashboard, a stack trace "
    "or an error dialog.\n"
    "Transcribe any error text, exception, hostname, timestamp or identifier "
    "you can read, exactly as it appears. Then say what it indicates.\n"
    "If the image is unreadable or is not what you expected, say so plainly. "
    "Never invent a hostname, an error code or a timestamp that is not "
    "visibly in the image - a fabricated identifier sends someone to the wrong "
    "machine."
)


def describe_image(image_bytes: bytes, question: Optional[str] = None,
                   timeout: Optional[int] = None) -> str:
    """Asks the vision model about an image.

    Goes to the same Ollama the chat already uses - `/api/generate` takes an
    `images` array of base64 strings for a multimodal model - so this adds no
    new service, port or model runtime to the deployment.
    """
    if not image_bytes:
        raise ValueError('no image received')
    if len(image_bytes) > MAX_IMAGE_BYTES:
        raise ValueError(f'image too large ({len(image_bytes)} bytes, limit {MAX_IMAGE_BYTES})')

    import requests
    payload = {
        'model': VISION_MODEL,
        'prompt': (question or 'What does this show? Transcribe any error text you can read.'),
        'system': _VISION_SYSTEM,
        'images': [base64.b64encode(image_bytes).decode('ascii')],
        'stream': False,
        'options': {'temperature': 0.1},
    }
    r = requests.post(f"{OLLAMA_HOST}/api/generate", json=payload,
                      timeout=timeout or VISION_TIMEOUT)
    if r.status_code != 200:
        raise RuntimeError(f"vision model returned HTTP {r.status_code}: {r.text[:200]}")
    return (r.json().get('response') or '').strip()
